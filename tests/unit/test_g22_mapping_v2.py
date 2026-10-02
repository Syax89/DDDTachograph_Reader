"""C3 (A-F4 / D2-001 / D2-002): G2.2 mapping across crypto, tree, export, pointer.

Covers the corrected card FID map, 0530 signature verification (incl. a
336-slot ECDSA control), the 0530 card/VU display-name collision, the
GNSS-accumulated-driving tree source, the shared export sections, the newest
cyclic pointer contract and the 0540 opaque byte EF. Real ECDSA crypto.
"""
import contextlib
import io
import struct

import pytest

from app.engine import TachoParser
from core.crypto.ef_signature import pair_ef_records
from core.decoders.card_g22 import parse_g22_gnss_accumulated_driving
from core.parser.deterministic import DeterministicParser
from core.registry.models import TachoResult, build_generations_tree
from core.registry.registry import DecoderRegistry, TagDecoder
from core.utils.report_format import (
    VERDICT_UNVERIFIED, VERDICT_VERIFIED, integrity_verdict, section_tables)
from scripts.rename_ddd_files import _integrity_status
from tests.unit.card_crypto import (
    ef_signature,
    g2_cert_records,
    g2_core_payloads,
    parse_bytes,
    signed_pairs,
    stap,
    trust_store,
    trusted_root_and_msca,
    v2_payloads,
)

TS = 1700000000


def _gnss_place(t=TS):
    coord = lambda v: int(v).to_bytes(3, "big", signed=True)
    return struct.pack(">I", t) + bytes([7]) + coord(45041) + coord(9125) + bytes([1])


# ── 0530: real 336-slot ECDSA control + byte tamper ────────────────────────

def _load_type_336():
    """A full 336-slot CardLoadTypeEntries: pointer(2) + 336×[ts(4)+type(1)]."""
    records = b"".join(
        struct.pack(">IB", TS, (i % 4) + 1) for i in range(336)
    )
    return b"\x00\x00" + records


def _full_v2_card(identity, load_type_payload, *, tamper_load_type=False):
    payloads = g2_core_payloads(v2=True)
    payloads.update(v2_payloads())
    payloads[0x0530] = load_type_payload
    payloads[0x0540] = b"opaque-vu-configuration-bytes"  # optional byte EF
    signed = {tag: value for tag, value in payloads.items() if tag != 0x0530}
    sent = bytearray(load_type_payload)
    if tamper_load_type:
        sent[6] ^= 0xFF  # first record's load-type byte (2-byte pointer + ts(4))
    return b"".join([
        g2_cert_records(identity["card_cert"], identity["msca_cert"]),
        signed_pairs(signed, identity["card_key"], 2),
        stap(0x0530, 0x02, bytes(sent)),
        stap(0x0530, 0x03, ef_signature(identity["card_key"], load_type_payload)),
    ])


def test_336_slot_load_type_full_v2_control_is_verified():
    erca_cert, identity = trusted_root_and_msca()
    load_type = _load_type_336()
    assert len(load_type) == 1682  # TCS_155 capacity
    with trust_store(g2_erca_cert=erca_cert) as certs_dir:
        result = parse_bytes(_full_v2_card(identity, load_type), certs_dir)[1]

    assert result["metadata"]["generation"] == "G2.2 (Smart V2)"
    assert result["metadata"]["integrity_check"] == "Verified"
    assert integrity_verdict(result) == VERDICT_VERIFIED
    assert result["ef_signature_verification"]["missing_core_efs"] == []
    assert len(result["load_type_entries"]) == 336
    # The optional 0540 opaque byte EF is signature-verified like any other EF.
    verified_tags = {r["tag"] for r in result["ef_signature_verification"]["ef_results"]
                     if r["status"] == "verified"}
    assert "0x0540" in verified_tags


def test_336_slot_load_type_tamper_fails_crypto_and_all_surfaces():
    erca_cert, identity = trusted_root_and_msca()
    load_type = _load_type_336()
    with trust_store(g2_erca_cert=erca_cert) as certs_dir:
        attack = parse_bytes(
            _full_v2_card(identity, load_type, tamper_load_type=True), certs_dir)[1]

    # The stale signature no longer covers the changed load-type byte.
    assert attack["load_type_entries"][0]["load_type"] == 1 ^ 0xFF
    assert attack["ef_signature_verification"]["failed"] >= 1
    assert not attack["metadata"]["integrity_check"].startswith("Verified")
    assert integrity_verdict(attack) == VERDICT_UNVERIFIED
    assert _integrity_status(attack) == "UNVERIFIED"

    pytest.importorskip("tkinter")
    from app.gui import TachoExplorer
    assert TachoExplorer._integrity_label(None, attack) != "All signatures verified"


# ── Newest cyclic pointer (5B records) ─────────────────────────────────────

def _parse_card_ef(tag, payload):
    return DeterministicParser().parse(stap(tag, 0x02, payload), is_vu=False)


def test_pointer_selects_newest_record_with_same_timestamps_different_codes():
    slots = struct.pack(">IBIB", TS, 1, TS, 2)  # two records, same ts, codes 1/2
    ptr0 = _parse_card_ef(0x0530, b"\x00\x00" + slots)
    ptr1 = _parse_card_ef(0x0530, b"\x00\x01" + slots)

    assert [r["load_type"] for r in ptr0["load_type_entries"]] == [1, 2]
    assert [r["record_index"] for r in ptr0["load_type_entries"]] == [0, 1]
    assert [r["is_newest"] for r in ptr0["load_type_entries"]] == [True, False]
    assert [r["load_type"] for r in ptr1["load_type_entries"]] == [1, 2]
    assert [r["is_newest"] for r in ptr1["load_type_entries"]] == [False, True]


def test_out_of_range_pointer_is_rejected_with_a_warning():
    slots = struct.pack(">IBIB", TS, 1, TS, 2)
    bad = _parse_card_ef(0x0530, b"\xff\xff" + slots)

    assert bad.get("load_type_entries") is None
    warning = bad["metadata"]["decoder_validation_warnings"][0]
    assert warning["code"] == "decoder_pointer_range_violation"
    assert warning["tag_id"] == "0x0530"


# ── GNSS accumulated driving: real VU RecordArray -> parser -> tree ────────

def _arr(rt, payload):
    return struct.pack(">BHH", rt, len(payload), 1) + payload


def test_vu_gnss_record_flows_to_gnss_tree_section(tmp_path):
    vu_record = struct.pack(">I", TS) + bytes(38) + _gnss_place() + (100).to_bytes(3, "big")
    data = (b"\x76\x31" + _arr(0x08, bytes(64))
            + b"\x76\x32" + _arr(0x16, vu_record) + _arr(0x08, bytes(64)))

    path = tmp_path / "vu_g22.ddd"
    path.write_bytes(data)
    result = TachoParser(str(path)).parse()

    assert len(result["gnss_ad_records"]) == 1
    tree = result["generations"]
    assert tree["Generation 2.2"]["GNSSAccumulatedDriving"] == result["gnss_ad_records"]
    assert tree["Generation 2"]["GNSSAccumulatedDriving"] == result["gnss_ad_records"]
    # The section is not mislabelled with the V2 Application_Identification tag.
    assert "DriverCardApplicationIdentificationV2" not in tree["Generation 2"]


def test_vu_gnss_decoder_emits_the_same_shape_directly():
    results = {}
    record = struct.pack(">I", TS) + _gnss_place(TS + 1) + (123456).to_bytes(3, "big")
    parse_g22_gnss_accumulated_driving(b"\x00\x00" + record, results)
    assert len(results["gnss_ad_records"]) == 1


# ── 0530 card/VU display-name collision (contextual lookup) ────────────────

def test_registry_resolves_0530_contextually_for_card_and_vu():
    registry = DecoderRegistry.instance()
    assert registry.get_decoder(0x0530, generation="G2.2", is_vu=False).name \
        == "G22_CardLoadTypeEntries"
    assert registry.get_decoder(0x0530, generation="G2.2", is_vu=True).name \
        == "VuPowerSupplyInterruptionData"


def test_tree_labels_card_0530_with_the_card_name(tmp_path):
    parser = TachoParser(str(tmp_path / "input.ddd"))
    result = TachoResult().to_dict()
    result["load_type_entries"] = [{"load_type": 1}]
    tree = build_generations_tree(result, parser.TAGS)

    assert "CardLoadTypeEntries" in tree["Generation 2.2"]
    assert "PowerSupplyInterruptionData" not in tree["Generation 2.2"]


def test_tree_keeps_custom_registry_driven_0530_name(tmp_path):
    registry = DecoderRegistry.instance()
    registry.register_decoder(TagDecoder(
        0x0530, "G22_RegistryLoadType", generation="G2.2", priority=1))

    parser = TachoParser(str(tmp_path / "input.ddd"))
    result = TachoResult().to_dict()
    result["load_type_entries"] = [{"load_type": 1}]
    tree = build_generations_tree(result, parser.TAGS)

    assert "RegistryLoadType" in tree["Generation 2.2"]


# ── 0540 opaque byte EF ────────────────────────────────────────────────────

def test_0540_recognized_as_a_signed_byte_ef():
    pairs = pair_ef_records([(0x0540, 0x02, b"opaque")], [(0x0540, 0x03, b"signature")])
    assert [p["tag"] for p in pairs] == [0x0540]
    assert pairs[0]["status"] == "paired"


# ── Shared export sections (CSV / Excel / PDF) ─────────────────────────────

def _export_synthetic():
    return {
        "metadata": {"filename": "s.ddd", "generation": "G2.2 (Smart V2)",
                     "is_vu": False, "file_size_bytes": 100, "coverage_pct": 100.0},
        "card_application_v2": {
            "length_of_following_data": 8, "no_border_crossing_records": 1,
            "no_load_unload_records": 1, "no_load_type_entry_records": 336,
            "vu_configuration_length_range": 3072},
        "place_auth_records": [{"timestamp": "2023-11-14T22:13:20+00:00", "authentication_status": 1}],
        "gnss_auth_records": [{"timestamp": "2023-11-14T22:13:20+00:00", "authentication_status": 1}],
        "load_type_entries": [{"timestamp": "2023-11-14T22:13:20+00:00", "load_type": 2}],
    }


def test_new_sections_are_in_the_shared_export_tables():
    labels = [row[0] for row in section_tables(_export_synthetic())]
    for expected in ("Card Application V2", "Places Authentication",
                     "GNSS Places Authentication", "Load Type Entries"):
        assert expected in labels


def test_new_sections_read_back_from_csv_excel_pdf(tmp_path):
    from app.export import ExportManager
    from openpyxl import load_workbook

    data = _export_synthetic()
    csv_path = tmp_path / "out.csv"
    xlsx_path = tmp_path / "out.xlsx"
    pdf_path = tmp_path / "out.pdf"

    ExportManager.export_to_csv(data, str(csv_path))
    ExportManager.export_to_excel(data, str(xlsx_path))
    ExportManager.export_to_pdf(data, str(pdf_path))

    csv_text = csv_path.read_text(encoding="utf-8-sig")
    assert "=== LOAD TYPE ENTRIES ===" in csv_text
    assert "Load Type" in csv_text
    assert "= PLACES AUTHENTICATION =" in csv_text

    wb = load_workbook(xlsx_path)
    sheets = set(wb.sheetnames)
    assert "Load Type Entries" in sheets
    assert "Card Application V2" in sheets
    assert "Places Authentication" in sheets

    import fitz
    with fitz.open(pdf_path) as pdf:
        pdf_text = "\n".join(page.get_text() for page in pdf)
    assert "Load Type Entries" in pdf_text
    assert "Card Application V2" in pdf_text


# ── Mock generator length field ────────────────────────────────────────────

def test_mock_g22_card_application_v2_length_field_is_8(tmp_path):
    from tests.integration import generate_mock_ddd as generator

    path = tmp_path / "mock_g22_card.ddd"
    with contextlib.redirect_stdout(io.StringIO()):
        generator.generate_g22_card(str(path))
    result = DeterministicParser().parse(path.read_bytes(), is_vu=False)

    assert result["card_application_v2"]["length_of_following_data"] == 8


# ── M3: every corrected V2 EF is really ECDSA-verified ─────────────────────

def _full_v2_payloads():
    payloads = g2_core_payloads(v2=True)
    payloads.update(v2_payloads())
    payloads[0x0540] = b"opaque-vu-configuration-bytes"
    return payloads


@pytest.mark.parametrize("tag", [0x0525, 0x0526, 0x0527, 0x0528, 0x0529, 0x0530, 0x0540])
def test_v2_ef_byte_tamper_fails_signature_and_all_surfaces(tag):
    erca_cert, identity = trusted_root_and_msca()
    payloads = _full_v2_payloads()
    genuine = payloads[tag]
    bad = bytearray(genuine)
    bad[len(bad) // 2] ^= 0xFF
    signed = {t: v for t, v in payloads.items() if t != tag}
    data = b"".join([
        g2_cert_records(identity["card_cert"], identity["msca_cert"]),
        signed_pairs(signed, identity["card_key"], 2),
        stap(tag, 0x02, bytes(bad)),
        stap(tag, 0x03, ef_signature(identity["card_key"], genuine)),
    ])
    with trust_store(g2_erca_cert=erca_cert) as certs_dir:
        attack = parse_bytes(data, certs_dir)[1]

    efv = attack["ef_signature_verification"]
    failed_tags = {r["tag"] for r in efv["ef_results"] if r["status"] == "failed"}
    assert f"0x{tag:04X}" in failed_tags
    assert efv["failed"] >= 1
    assert attack["metadata"]["integrity_check"] == "Unverified (EF Signature Mismatch)"
    assert integrity_verdict(attack) == VERDICT_UNVERIFIED
    assert _integrity_status(attack) == "UNVERIFIED"


def test_full_v2_signed_control_is_verified_before_tampering():
    """Positive control: the same fixture, untampered, is actually Verified and
    every corrected V2 EF is in the verified set."""
    erca_cert, identity = trusted_root_and_msca()
    payloads = _full_v2_payloads()
    data = (g2_cert_records(identity["card_cert"], identity["msca_cert"])
            + signed_pairs(payloads, identity["card_key"], 2))
    with trust_store(g2_erca_cert=erca_cert) as certs_dir:
        result = parse_bytes(data, certs_dir)[1]

    assert result["metadata"]["integrity_check"] == "Verified"
    assert integrity_verdict(result) == VERDICT_VERIFIED
    verified = {r["tag"] for r in result["ef_signature_verification"]["ef_results"]
                if r["status"] == "verified"}
    for tag in (0x0525, 0x0526, 0x0527, 0x0528, 0x0529, 0x0530, 0x0540):
        assert f"0x{tag:04X}" in verified


# ── F2: newest pointer retained for the 0528/0529 card records ─────────────

def _0528_record(ts=TS):
    return bytes([0x1A, 0x0D]) + _gnss_place(ts) + (100).to_bytes(3, "big")


def _0529_record(ts=TS):
    return struct.pack(">IB", ts, 1) + _gnss_place(ts) + (100).to_bytes(3, "big")


@pytest.mark.parametrize(("tag", "key", "make_record"), [
    (0x0528, "border_crossings", _0528_record),
    (0x0529, "load_unload_records", _0529_record),
])
def test_0528_0529_pointer_selects_newest_record(tag, key, make_record):
    slots = make_record() * 2  # two records with identical timestamps
    ptr0 = _parse_card_ef(tag, b"\x00\x00" + slots)
    ptr1 = _parse_card_ef(tag, b"\x00\x01" + slots)

    assert [r["record_index"] for r in ptr0[key]] == [0, 1]
    assert [r["is_newest"] for r in ptr0[key]] == [True, False]
    assert [r["is_newest"] for r in ptr1[key]] == [False, True]
    assert ptr0[key] != ptr1[key]  # pointer 0 vs 1 must be distinguishable


@pytest.mark.parametrize("tag", [0x0528, 0x0529])
def test_0528_0529_out_of_range_pointer_is_rejected(tag):
    slots = (_0528_record() if tag == 0x0528 else _0529_record()) * 2
    bad = _parse_card_ef(tag, b"\xff\xff" + slots)

    key = "border_crossings" if tag == 0x0528 else "load_unload_records"
    assert not bad.get(key)
    warning = bad["metadata"]["decoder_validation_warnings"][0]
    assert warning["code"] == "decoder_pointer_range_violation"


# ── F3: VU RecordArray context in the Gen 2.2 tree ─────────────────────────

def _vu_record_array_file():
    gnss = struct.pack(">I", TS) + bytes(38) + _gnss_place() + (100).to_bytes(3, "big")
    border = bytes(38) + bytes([0x1A, 0x0D]) + _gnss_place() + (100).to_bytes(3, "big")
    unload = struct.pack(">IB", TS, 1) + bytes(38) + _gnss_place() + (100).to_bytes(3, "big")
    return (b"\x76\x31" + _arr(0x08, bytes(64))
            + b"\x76\x32" + _arr(0x16, gnss) + _arr(0x22, border) + _arr(0x23, unload)
            + _arr(0x08, bytes(64)))


def test_vu_tree_uses_vu_section_names_and_keeps_gnss(tmp_path):
    path = tmp_path / "vu_g22.ddd"
    path.write_bytes(_vu_record_array_file())
    result = TachoParser(str(path)).parse()

    g22 = result["generations"]["Generation 2.2"]
    # VU provenance: source-agnostic section names, not the card-only FIDs.
    assert "BorderCrossings" in g22 and "CardBorderCrossings" not in g22
    assert "LoadUnloadOperations" in g22 and "CardLoadUnloadOperations" not in g22
    assert g22["GNSSAccumulatedDriving"] == result["gnss_ad_records"]
    assert g22["BorderCrossings"] == result["border_crossings"]
    assert g22["LoadUnloadOperations"] == result["load_unload_records"]
    # VU row identity/values preserved, and VU RecordArray fields are untouched
    # by the card pointer contract.
    assert result["border_crossings"][0]["name"] == "VuBorderCrossingRecord"
    assert result["load_unload_records"][0]["name"] == "VuLoadUnloadRecord"
    assert result["border_crossings"][0]["odometer_km"] == 100
    assert "record_index" not in result["border_crossings"][0]
    assert "is_newest" not in result["load_unload_records"][0]


def test_card_tree_uses_card_section_names(tmp_path):
    parser = TachoParser(str(tmp_path / "input.ddd"))
    result = TachoResult().to_dict()
    result["metadata"]["is_vu"] = False
    result["border_crossings"] = [{"nation_from": "I"}]
    result["load_unload_records"] = [{"operation": "LOAD"}]
    tree = build_generations_tree(result, parser.TAGS)

    g22 = tree["Generation 2.2"]
    assert "CardBorderCrossings" in g22 and "BorderCrossings" not in g22
    assert "CardLoadUnloadOperations" in g22 and "LoadUnloadOperations" not in g22


def test_card_tree_keeps_custom_registry_0528_name(tmp_path):
    registry = DecoderRegistry.instance()
    registry.register_decoder(TagDecoder(
        0x0528, "G22_RegistryBorderCrossings", generation="G2.2", priority=1))

    parser = TachoParser(str(tmp_path / "input.ddd"))
    result = TachoResult().to_dict()
    result["metadata"]["is_vu"] = False
    result["border_crossings"] = [{"nation_from": "I"}]
    tree = build_generations_tree(result, parser.TAGS)

    assert "RegistryBorderCrossings" in tree["Generation 2.2"]
